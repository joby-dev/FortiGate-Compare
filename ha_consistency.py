"""HA configuration consistency analysis for FortiGate exports."""

from __future__ import annotations

import difflib
import html
import re
import shlex
from dataclasses import dataclass, field
from io import BytesIO, StringIO

from openpyxl import Workbook


SCOPE_CATEGORIES = {
    "firewall policy": "Firewall Policies",
    "firewall proxy-policy": "Proxy Policies",
    "firewall local-in-policy": "Local-In Policies",
    "firewall address": "Address Objects",
    "firewall addrgrp": "Address Groups",
    "firewall service custom": "Service Objects",
    "firewall service group": "Service Groups",
    "firewall vip": "VIPs",
    "firewall vipgrp": "VIPs",
    "router static": "Static Routes",
    "system sdwan": "SD-WAN",
    "system virtual-wan-link": "SD-WAN",
    "firewall profile-protocol-options": "Security Profiles",
    "firewall profile-group": "Security Profiles",
    "firewall ips sensor": "Security Profiles",
    "firewall antivirus profile": "Security Profiles",
    "firewall voip profile": "Security Profiles",
    "firewall ssl-ssh-profile": "SSL Inspection Profiles",
    "webfilter profile": "Web Filter Profiles",
    "dnsfilter profile": "DNS Filter Profiles",
    "application list": "Application Control Profiles",
}

POLICY_CATEGORIES = {"Firewall Policies", "Proxy Policies", "Local-In Policies"}
POLICY_MATCH_MODES = ("Policy ID Match", "Functional Match")
PROFILE_CATEGORIES = {
    "Security Profiles",
    "SSL Inspection Profiles",
    "Web Filter Profiles",
    "DNS Filter Profiles",
    "Application Control Profiles",
}

_IGNORED_LINE = re.compile(
    r"^\s*(?:#\s*)?(?:set\s+)?(?:config-version|build(?:no|number|version)?|"
    r"conf_file_ver|revision(?:-number)?|timestamp|last-modified|serial-number|"
    r"firmware-version|firmware-build)\b",
    re.IGNORECASE,
)
_COMMENT_LINE = re.compile(r"^(?:#|!|//)")
_COSMETIC_SETTING = re.compile(r"^set\s+(?:comments?|uuid(?:-index)?)\b", re.IGNORECASE)
_REFERENCE_FIELDS = {
    "srcaddr",
    "dstaddr",
    "srcaddr6",
    "dstaddr6",
    "service",
}
_PROFILE_FIELDS = {
    "av-profile",
    "webfilter-profile",
    "ips-sensor",
    "application-list",
    "ssl-ssh-profile",
    "dnsfilter-profile",
    "profile-protocol-options",
    "voip-profile",
    "file-filter-profile",
    "emailfilter-profile",
    "dlp-sensor",
    "icap-profile",
    "waf-profile",
    "virtual-patch-profile",
}
_POLICY_FUNCTION_FIELDS = {
    "srcintf",
    "dstintf",
    "srcaddr",
    "dstaddr",
    "service",
    "action",
    "schedule",
    "nat",
} | _PROFILE_FIELDS
_BUILTIN_REFERENCES = {"all", "none", "any", "always", "default"}
_POLICY_FINGERPRINT_SIMILARITY_THRESHOLD = 0.9


@dataclass(frozen=True)
class HAObject:
    key: str
    category: str
    object_id: str
    label: str
    context: str
    lines: tuple[str, ...]
    cosmetic_lines: tuple[str, ...] = ()


@dataclass(frozen=True)
class HAChange:
    category: str
    object_id: str
    label: str
    status: str
    before_lines: tuple[str, ...]
    after_lines: tuple[str, ...]
    risk: str = ""
    risk_reason: str = ""
    impact: str = ""
    recommendation: str = ""
    before_object_id: str = ""
    after_object_id: str = ""


@dataclass
class HAComparison:
    before_name: str
    after_name: str
    changes: list[HAChange]
    dependencies: list[dict[str, str]] = field(default_factory=list)
    risks: list[dict[str, str]] = field(default_factory=list)
    cosmetic_changes: list[HAChange] = field(default_factory=list)

    @property
    def matching_count(self) -> int:
        return sum(change.status == "Matching" for change in self.changes)

    @property
    def different_count(self) -> int:
        return sum(change.status == "Modified" for change in self.changes)

    @property
    def missing_count(self) -> int:
        return sum(change.status.startswith("Missing in") for change in self.changes)

    @property
    def compared_count(self) -> int:
        return len(self.changes)

    @property
    def exact_policy_count(self) -> int:
        return sum(
            change.category in POLICY_CATEGORIES and change.status == "Matching"
            for change in self.changes
        )

    @property
    def equivalent_policy_count(self) -> int:
        return sum(
            change.category in POLICY_CATEGORIES and change.status == "Equivalent"
            for change in self.changes
        )

    @property
    def modified_policy_count(self) -> int:
        return sum(
            change.category in POLICY_CATEGORIES and change.status == "Modified"
            for change in self.changes
        )

    @property
    def missing_policy_count(self) -> int:
        return sum(
            change.category in POLICY_CATEGORIES and change.status.startswith("Missing in")
            for change in self.changes
        )

    @property
    def cosmetic_count(self) -> int:
        return len(self.cosmetic_changes)

    @property
    def health_score(self) -> int:
        total_findings = len(self.changes)
        if not total_findings:
            return 100
        equivalent_count = sum(change.status == "Equivalent" for change in self.changes)
        return round((self.matching_count + equivalent_count) * 100 / total_findings)


def _tokens(line: str) -> list[str]:
    try:
        return shlex.split(line, posix=True)
    except ValueError:
        return line.split()


def _config_category(config_path: str) -> str | None:
    path = re.sub(r"\s+", " ", config_path.strip().lower())
    if path in SCOPE_CATEGORIES:
        return SCOPE_CATEGORIES[path]
    return None


def _normalize_line(line: str) -> str:
    pieces = re.split(r'("(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\')', line.strip())
    return "".join(
        piece if index % 2 else re.sub(r"\s+", " ", piece)
        for index, piece in enumerate(pieces)
    )


def _parse_objects(
    config_text: str,
) -> dict[str, HAObject]:
    """Index relevant edit blocks without splitting the full file into lines."""
    if not isinstance(config_text, str) or not config_text.strip():
        raise ValueError("The uploaded configuration is empty or is not valid text.")

    objects: dict[str, HAObject] = {}
    config_stack: list[str] = []
    edit_stack: list[tuple[int, str]] = []
    active_category: str | None = None
    active_id = ""
    active_key = ""
    active_depth = -1
    active_context = ""
    active_lines: list[str] = []
    section_scope_depth = -1
    section_scope_category: str | None = None

    def finish_object() -> None:
        nonlocal active_category, active_id, active_key, active_depth, active_context, active_lines
        if active_category is None:
            return
        lines = []
        cosmetic_lines = []
        for raw in active_lines:
            stripped = raw.strip()
            if not stripped or _IGNORED_LINE.match(stripped):
                continue
            normalized = _normalize_line(raw)
            if _COMMENT_LINE.match(stripped) or _COSMETIC_SETTING.match(stripped):
                cosmetic_lines.append(normalized)
            else:
                lines.append(normalized)
        lines = tuple(lines)
        cosmetic_lines = tuple(cosmetic_lines)
        if active_category == "SD-WAN":
            object_id = "SD-WAN configuration"
            key = f"{active_category}::{object_id}"
        else:
            object_id = active_id
            key = f"{active_category}::{active_context}::{object_id}"
        label = object_id
        for line in lines:
            tokens = _tokens(line)
            if len(tokens) > 2 and tokens[0].lower() == "set" and tokens[1].lower() == "name":
                label = " ".join(tokens[2:])
                break
        obj = HAObject(
            key,
            active_category,
            object_id,
            label,
            active_context,
            lines,
            cosmetic_lines,
        )
        objects[key] = obj
        active_category = None
        active_id = ""
        active_key = ""
        active_depth = -1
        active_context = ""
        active_lines = []

    for raw_line in StringIO(config_text):
        stripped = raw_line.strip()
        if not stripped:
            continue
        tokens = _tokens(stripped)
        if not tokens:
            continue
        command = tokens[0].lower()
        current_depth = len(config_stack)

        if command == "config":
            config_path = " ".join(tokens[1:])
            category = _config_category(config_path)
            if active_category is not None:
                active_lines.append(stripped)
            config_stack.append(config_path)
            if active_category is None and category == "SD-WAN":
                context = " / ".join(value for _, value in edit_stack)
                active_category = category
                active_id = "SD-WAN configuration"
                active_key = f"{category}::{context}::{active_id}"
                active_depth = len(config_stack)
                active_context = context
                section_scope_depth = active_depth
                section_scope_category = category
                active_lines = [stripped]
            continue

        if command == "edit":
            edit_id = " ".join(tokens[1:]) if len(tokens) > 1 else "<unnamed>"
            if active_category is None and current_depth:
                category = _config_category(config_stack[-1])
                if category and category != "SD-WAN":
                    active_category = category
                    active_id = edit_id.strip('"\'')
                    context = " / ".join(value for depth, value in edit_stack if depth < current_depth)
                    active_key = f"{category}::{context}::{active_id}"
                    active_depth = current_depth
                    active_context = context
                    active_lines = [stripped]
            elif active_category is not None:
                active_lines.append(stripped)
            edit_stack.append((current_depth, edit_id.strip('"\'')))
            continue

        if command == "next":
            if active_category is not None:
                active_lines.append(stripped)
            matching_edit = next(
                (index for index in range(len(edit_stack) - 1, -1, -1) if edit_stack[index][0] == current_depth),
                None,
            )
            if matching_edit is not None:
                edit_stack.pop(matching_edit)
            if active_category is not None and active_depth == current_depth:
                finish_object()
            continue

        if command == "end":
            if active_category is not None:
                active_lines.append(stripped)
            if config_stack:
                if section_scope_depth == current_depth and section_scope_category == "SD-WAN":
                    finish_object()
                    section_scope_depth = -1
                    section_scope_category = None
                config_stack.pop()
                edit_stack[:] = [item for item in edit_stack if item[0] < current_depth]
            continue

        if active_category is not None:
            active_lines.append(stripped)

    if active_category is not None:
        finish_object()
    return objects


def _object_values(obj: HAObject, fields: set[str]) -> dict[str, list[str]]:
    values: dict[str, list[str]] = {}
    for line in obj.lines:
        tokens = _tokens(line)
        if len(tokens) >= 3 and tokens[0].lower() == "set" and tokens[1].lower() in fields:
            values[tokens[1].lower()] = tokens[2:]
    return values


def _policy_fingerprint(obj: HAObject) -> tuple[tuple[str, tuple[str, ...]], ...]:
    values = _object_values(obj, _POLICY_FUNCTION_FIELDS)
    return tuple(
        (field_name, tuple(values.get(field_name, ())))
        for field_name in sorted(_POLICY_FUNCTION_FIELDS)
    )


def _policy_fingerprint_similarity(before: HAObject, after: HAObject) -> float:
    before_fingerprint = _policy_fingerprint(before)
    after_fingerprint = _policy_fingerprint(after)
    if not before_fingerprint:
        return 1.0
    matching_fields = sum(
        before_value == after_value
        for (_, before_value), (_, after_value) in zip(
            before_fingerprint, after_fingerprint
        )
    )
    return matching_fields / len(before_fingerprint)


def _pair_policies(
    before_objects: dict[str, HAObject],
    after_objects: dict[str, HAObject],
    mode: str,
) -> list[tuple[HAObject | None, HAObject | None, str]]:
    before_policies = [obj for obj in before_objects.values() if obj.category in POLICY_CATEGORIES]
    after_policies = [obj for obj in after_objects.values() if obj.category in POLICY_CATEGORIES]
    before_unmatched = {obj.key: obj for obj in before_policies}
    after_unmatched = {obj.key: obj for obj in after_policies}
    pairs: list[tuple[HAObject | None, HAObject | None, str]] = []

    if mode == "Functional Match":
        after_by_fingerprint: dict[tuple[str, str, tuple[tuple[str, tuple[str, ...]], ...]], list[HAObject]] = {}
        for after in after_unmatched.values():
            identity = (after.category, after.context, _policy_fingerprint(after))
            after_by_fingerprint.setdefault(identity, []).append(after)
        for candidates in after_by_fingerprint.values():
            candidates.sort(key=lambda obj: obj.object_id.casefold())
        for before in list(before_unmatched.values()):
            identity = (before.category, before.context, _policy_fingerprint(before))
            candidates = after_by_fingerprint.get(identity, [])
            if not candidates:
                continue
            after = candidates.pop(0)
            del before_unmatched[before.key]
            del after_unmatched[after.key]
            status = (
                "Matching"
                if before.object_id == after.object_id and before.lines == after.lines
                else "Equivalent"
            )
            pairs.append((before, after, status))

        for before in list(before_unmatched.values()):
            if not before.label.strip():
                continue
            candidates = sorted(
                (
                    (_policy_fingerprint_similarity(before, after), after)
                    for after in after_unmatched.values()
                    if after.category == before.category
                    and after.context == before.context
                    and after.label.strip().casefold() == before.label.strip().casefold()
                ),
                key=lambda item: (-item[0], item[1].object_id.casefold()),
            )
            if not candidates or candidates[0][0] < _POLICY_FINGERPRINT_SIMILARITY_THRESHOLD:
                continue
            after = candidates[0][1]
            del before_unmatched[before.key]
            del after_unmatched[after.key]
            pairs.append((before, after, "Equivalent"))

    after_by_identity = {
        (obj.category, obj.context, obj.object_id): obj
        for obj in after_unmatched.values()
    }
    for before in list(before_unmatched.values()):
        identity = (before.category, before.context, before.object_id)
        after = after_by_identity.get(identity)
        if after is None:
            continue
        del before_unmatched[before.key]
        del after_unmatched[after.key]
        status = "Matching" if before.lines == after.lines else "Modified"
        pairs.append((before, after, status))

    pairs.extend((before, None, "Missing in B") for before in before_unmatched.values())
    pairs.extend((None, after, "Missing in A") for after in after_unmatched.values())
    return pairs


def _risk_for_change(change: HAChange) -> tuple[str, str]:
    matcher = difflib.SequenceMatcher(
        None, change.before_lines, change.after_lines, autojunk=False
    )
    changed_lines = [
        line
        for tag, before_start, before_end, after_start, after_end in matcher.get_opcodes()
        if tag != "equal"
        for line in (
            change.before_lines[before_start:before_end]
            + change.after_lines[after_start:after_end]
        )
    ]
    combined = "\n".join(changed_lines).casefold()
    if change.category in POLICY_CATEGORIES and change.status.startswith("Missing in"):
        return "High", "A firewall policy is missing on one peer."
    if change.category in PROFILE_CATEGORIES and change.status.startswith("Missing in"):
        return "High", "A security profile is missing on one peer."
    if change.status == "Modified" and re.search(r"\bset\s+(?:nat|ippool|poolname)\b", combined):
        return "High", "NAT or IP pool behavior differs between peers."
    if change.status == "Modified" and re.search(
        r"\bset\s+(?:ips-sensor|webfilter-profile|av-profile|application-list|"
        r"ssl-ssh-profile|dnsfilter-profile|profile-protocol-options|voip-profile)\b",
        combined,
    ):
        return "High", "Security inspection profiles differ between peers."
    if change.status == "Modified" and re.search(r"\bset\s+(?:service)\b", combined):
        return "Medium", "A service reference differs between peers."
    if change.status == "Modified" and re.search(r"\bset\s+(?:srcaddr|dstaddr|srcaddr6|dstaddr6)\b", combined):
        return "Medium", "An address reference differs between peers."
    if change.status == "Modified" and re.search(r"\bset\s+comments?\b", combined):
        return "Low", "Only configuration comments differ."
    if change.status.startswith("Missing in"):
        return "Medium", "A configuration object exists on only one peer."
    if change.status == "Modified":
        return "Medium", "Functional settings differ between peers."
    return "", ""


def exact_differences(change: HAChange) -> list[dict[str, str]]:
    """Return setting-level values for a paired HA object difference."""
    if change.status == "Equivalent":
        return []
    if change.status.startswith("Missing in"):
        missing_side = "Firewall A" if change.status == "Missing in A" else "Firewall B"
        present_side = "Firewall B" if missing_side == "Firewall A" else "Firewall A"
        return [
            {
                "Setting": "Object",
                "Firewall A": "Not present" if missing_side == "Firewall A" else "Present",
                "Firewall B": "Not present" if missing_side == "Firewall B" else "Present",
                "Difference": f"Object exists only on {present_side}.",
            }
        ]

    matcher = difflib.SequenceMatcher(
        None, change.before_lines, change.after_lines, autojunk=False
    )
    differences: list[dict[str, str]] = []
    for tag, before_start, before_end, after_start, after_end in matcher.get_opcodes():
        if tag == "equal":
            continue
        before_chunk = change.before_lines[before_start:before_end]
        after_chunk = change.after_lines[after_start:after_end]
        paired = max(len(before_chunk), len(after_chunk))
        for offset in range(paired):
            before_line = before_chunk[offset] if offset < len(before_chunk) else ""
            after_line = after_chunk[offset] if offset < len(after_chunk) else ""
            before_tokens = _tokens(before_line)
            after_tokens = _tokens(after_line)
            before_field = (
                before_tokens[1]
                if len(before_tokens) > 1 and before_tokens[0].lower() in {"set", "unset"}
                else "Configuration"
            )
            after_field = (
                after_tokens[1]
                if len(after_tokens) > 1 and after_tokens[0].lower() in {"set", "unset"}
                else "Configuration"
            )
            before_value = " ".join(before_tokens[2:]) if before_field != "Configuration" else before_line
            after_value = " ".join(after_tokens[2:]) if after_field != "Configuration" else after_line
            setting = after_field if after_field != "Configuration" else before_field
            if before_field != after_field:
                setting = "Configuration"
            if before_value == after_value:
                continue
            differences.append(
                {
                    "Setting": setting,
                    "Firewall A": before_value or "Not set",
                    "Firewall B": after_value or "Not set",
                    "Difference": f"{before_value or 'Not set'} -> {after_value or 'Not set'}",
                }
            )
    return differences


def _review_guidance(
    change: HAChange, risk: str
) -> tuple[str, str]:
    if change.status == "Equivalent":
        return (
            "The policies have matching functional fingerprints despite different policy IDs.",
            "No action is required unless policy IDs must also be standardized.",
        )
    if change.status.startswith("Missing in"):
        present_side = "Firewall B" if change.status == "Missing in A" else "Firewall A"
        if change.category in POLICY_CATEGORIES:
            return (
                f"Traffic matching this policy may be handled differently when {present_side} processes traffic.",
                "Confirm the policy is intended on both peers; align its settings and sequence.",
            )
        if change.category in PROFILE_CATEGORIES:
            return (
                f"Policies may receive different inspection when {present_side} processes traffic.",
                "Align the profile object and verify policy references on both peers.",
            )
        return (
            f"Policies on {present_side} may resolve this object differently or fail to match it.",
            "Confirm the object is required and synchronize it or update dependent references.",
        )

    fields = {row["Setting"].casefold() for row in exact_differences(change)}
    if "nat" in fields or "ippool" in fields or "poolname" in fields:
        return (
            "Traffic may be translated differently depending on which HA peer processes it.",
            "Align NAT mode, IP pools, and related policy settings across both peers.",
        )
    if fields & {"ips-sensor", "webfilter-profile", "av-profile", "application-list", "ssl-ssh-profile", "dnsfilter-profile", "profile-protocol-options", "voip-profile"}:
        if "ips-sensor" in fields:
            impact = "Traffic may receive different IPS inspection depending on the active processing node."
            action = "Align IPS profile configuration on both HA peers."
        else:
            impact = "Traffic may receive different security inspection depending on the active processing node."
            action = "Align the security profile assignment and referenced profile on both HA peers."
        return impact, action
    if "service" in fields:
        return (
            "Traffic may be permitted or denied differently depending on which peer processes it.",
            "Compare the service objects and align the policy service selection on both peers.",
        )
    if fields & {"srcaddr", "dstaddr", "srcaddr6", "dstaddr6"}:
        return (
            "Different address matches may change which traffic the policy handles.",
            "Verify address and address-group membership, then align the policy references.",
        )
    if change.category in POLICY_CATEGORIES:
        return (
            "Policy behavior may differ between peers for matching traffic.",
            "Review the changed settings and synchronize the intended policy behavior.",
        )
    if risk == "Low":
        return (
            "No functional traffic impact is expected from this cosmetic difference.",
            "No action is required unless the comment or metadata must be standardized.",
        )
    return (
        "Object behavior may differ when either HA peer handles traffic.",
        "Review the exact setting delta and align both peers to the approved configuration.",
    )


def _dependency_findings(
    before_objects: dict[str, HAObject], after_objects: dict[str, HAObject]
) -> list[dict[str, str]]:
    findings: list[dict[str, str]] = []
    for side, objects in (("Firewall A", before_objects), ("Firewall B", after_objects)):
        available: dict[tuple[str, str], set[str]] = {}
        for obj in objects.values():
            available.setdefault((obj.category, obj.context), set()).update(
                {obj.object_id.casefold(), obj.label.strip('"\'').casefold()}
            )

        for policy in objects.values():
            if policy.category not in POLICY_CATEGORIES:
                continue
            address_names = available.get(("Address Objects", policy.context), set()) | available.get(("Address Groups", policy.context), set())
            service_names = available.get(("Service Objects", policy.context), set()) | available.get(("Service Groups", policy.context), set())
            vip_names = available.get(("VIPs", policy.context), set())
            profile_names = {
                name
                for category in PROFILE_CATEGORIES
                for name in available.get((category, policy.context), set())
            }
            values = _object_values(policy, _REFERENCE_FIELDS)
            for field_name, names in values.items():
                for name in names:
                    normalized = name.casefold()
                    if normalized in _BUILTIN_REFERENCES:
                        continue
                    candidates: list[tuple[str, set[str]]]
                    if field_name == "service":
                        candidates = [("Service Object", service_names)]
                    elif field_name.startswith("dstaddr"):
                        candidates = [("Address Object", address_names | vip_names)]
                    else:
                        candidates = [("Address Object", address_names)]
                    for reference_type, available_names in candidates:
                        if normalized not in available_names:
                            findings.append(
                                {
                                    "Firewall": side,
                                    "Policy": f"{policy.category} {policy.label}",
                                    "Missing Reference": name,
                                    "Reference Type": reference_type,
                                    "Severity": "Medium",
                                }
                            )
            profile_values = _object_values(policy, _PROFILE_FIELDS)
            for field_name, names in profile_values.items():
                for name in names:
                    normalized = name.casefold()
                    if normalized in _BUILTIN_REFERENCES or normalized in profile_names:
                        continue
                    findings.append(
                        {
                            "Firewall": side,
                            "Policy": f"{policy.category} {policy.label}",
                            "Missing Reference": name,
                            "Reference Type": f"{field_name} security profile",
                            "Severity": "High",
                        }
                    )

    unique = {
        (row["Firewall"], row["Policy"], row["Missing Reference"], row["Reference Type"]): row
        for row in findings
    }
    return list(unique.values())


def compare_ha_configurations(
    before_text: str,
    after_text: str,
    before_name: str = "Firewall A",
    after_name: str = "Firewall B",
    policy_match_mode: str = "Policy ID Match",
) -> HAComparison:
    """Compare in-scope HA objects while excluding device-local sections."""
    if policy_match_mode not in POLICY_MATCH_MODES:
        raise ValueError(f"Unsupported policy comparison mode: {policy_match_mode}")
    before_objects = _parse_objects(before_text)
    after_objects = _parse_objects(after_text)
    policy_pairs = _pair_policies(before_objects, after_objects, policy_match_mode)
    changes: list[HAChange] = []
    cosmetic_changes: list[HAChange] = []
    for key in sorted(before_objects.keys() | after_objects.keys()):
        before = before_objects.get(key)
        after = after_objects.get(key)
        exemplar = before or after
        assert exemplar is not None
        if exemplar.category in POLICY_CATEGORIES:
            continue
        if before is None:
            status = "Missing in A"
            old_lines: tuple[str, ...] = ()
            new_lines = after.lines if after else ()
        elif after is None:
            status = "Missing in B"
            old_lines = before.lines
            new_lines = ()
        elif before.lines == after.lines:
            status = "Matching"
            old_lines = before.lines
            new_lines = after.lines
        else:
            status = "Modified"
            old_lines = before.lines
            new_lines = after.lines
        provisional = HAChange(
            exemplar.category,
            exemplar.object_id,
            exemplar.label,
            status,
            old_lines,
            new_lines,
        )
        risk, reason = _risk_for_change(provisional)
        impact, recommendation = _review_guidance(provisional, risk)
        changes.append(
            HAChange(
                provisional.category,
                provisional.object_id,
                provisional.label,
                provisional.status,
                provisional.before_lines,
                provisional.after_lines,
                risk,
                reason,
                impact,
                recommendation,
            )
        )
        if (
            before is not None
            and after is not None
            and before.cosmetic_lines != after.cosmetic_lines
        ):
            cosmetic_changes.append(
                HAChange(
                    exemplar.category,
                    exemplar.object_id,
                    exemplar.label,
                    "Cosmetic difference",
                    before.cosmetic_lines,
                    after.cosmetic_lines,
                    "Informational",
                    "Comments or UUID metadata differ; functional health is unaffected.",
                    "No functional traffic impact is expected from this cosmetic difference.",
                    "No action is required unless the comment or UUID metadata must be standardized.",
                )
            )

    for before, after, status in policy_pairs:
        exemplar = before or after
        assert exemplar is not None
        old_lines = before.lines if before is not None else ()
        new_lines = after.lines if after is not None else ()
        provisional = HAChange(
            exemplar.category,
            before.object_id if before is not None else after.object_id,
            exemplar.label,
            status,
            old_lines,
            new_lines,
            before_object_id=before.object_id if before is not None else "",
            after_object_id=after.object_id if after is not None else "",
        )
        risk, reason = _risk_for_change(provisional)
        impact, recommendation = _review_guidance(provisional, risk)
        changes.append(
            HAChange(
                provisional.category,
                provisional.object_id,
                provisional.label,
                provisional.status,
                provisional.before_lines,
                provisional.after_lines,
                risk,
                reason,
                impact,
                recommendation,
                provisional.before_object_id,
                provisional.after_object_id,
            )
        )
        if (
            before is not None
            and after is not None
            and before.cosmetic_lines != after.cosmetic_lines
        ):
            cosmetic_changes.append(
                HAChange(
                    exemplar.category,
                    before.object_id,
                    exemplar.label,
                    "Cosmetic difference",
                    before.cosmetic_lines,
                    after.cosmetic_lines,
                    "Informational",
                    "Comments or UUID metadata differ; functional health is unaffected.",
                    "No functional traffic impact is expected from this cosmetic difference.",
                    "No action is required unless the comment or UUID metadata must be standardized.",
                    before.object_id,
                    after.object_id,
                )
            )

    dependencies = _dependency_findings(before_objects, after_objects)
    risks = [
        {
            "Risk": change.risk,
            "Object": f"{change.category}: {change.label}",
            "Finding": change.risk_reason,
        }
        for change in changes
        if change.risk
    ]
    risks.extend(
        {
            "Risk": item["Severity"],
            "Object": item["Policy"],
            "Finding": f"{item['Firewall']} references missing {item['Reference Type'].lower()} '{item['Missing Reference']}'.",
        }
        for item in dependencies
    )
    return HAComparison(
        before_name,
        after_name,
        changes,
        dependencies,
        risks,
        cosmetic_changes,
    )


def render_ha_object_html(change: HAChange) -> str:
    """Render one object as an aligned side-by-side configuration diff."""
    before_lines = list(change.before_lines) or ["<OBJECT MISSING>"]
    after_lines = list(change.after_lines) or ["<OBJECT MISSING>"]
    table = difflib.HtmlDiff(wrapcolumn=120).make_table(
        before_lines,
        after_lines,
        fromdesc="Pre-Change",
        todesc="Post-Change",
        context=True,
        numlines=0,
    )
    return "".join(
        (
            "<!doctype html><html><head><meta charset='utf-8'><style>",
            "*{box-sizing:border-box}body{margin:0;background:#FFFFFF;color:#20252a;",
            "font:12px/1.5 'Segoe UI',Consolas,monospace}.wrap{height:100vh;overflow:auto}",
            "table.diff{width:100%;border-collapse:collapse;table-layout:fixed}",
            ".diff th{position:sticky;top:0;background:#D6EEFF;color:#20252a;padding:8px;text-align:left}",
            ".diff td{padding:3px 7px;border-bottom:1px solid #B8DFF5;vertical-align:top}",
            ".diff_header{background:#EAF6FF;color:#59636d}.diff_add{background:#D4F8D4;color:#173d22}",
            ".diff_sub{background:#FFD6D6;color:#5d1e1e}.diff_chg{background:#FFF1B8;color:#443900}",
            "td:nth-child(2),td:nth-child(5){width:45%;overflow-wrap:anywhere}",
            "a{color:#1769aa} </style></head><body><div class='wrap'>",
            table,
            "</div></body></html>",
        )
    )


def _append_html_finding(parts: list[str], change: HAChange) -> None:
    parts.append(
        f"<section><h3>{html.escape(change.category)} · {html.escape(change.label)}</h3>"
        f"<p><b>Status:</b> {html.escape(change.status)} &nbsp; "
        f"<b>Risk:</b> {html.escape(change.risk or 'Informational')}</p>"
        "<table><tr><th>Setting</th><th>Firewall A</th><th>Firewall B</th><th>Exact Difference</th></tr>"
    )
    for item in exact_differences(change):
        parts.append(
            "<tr>"
            + "".join(f"<td>{html.escape(item[key])}</td>" for key in ("Setting", "Firewall A", "Firewall B", "Difference"))
            + "</tr>"
        )
    parts.append(
        "</table>"
        f"<p><b>Impact Assessment:</b> {html.escape(change.impact or change.risk_reason)}</p>"
        f"<p><b>Recommended Action:</b> {html.escape(change.recommendation)}</p>"
        "</section>"
    )


def build_ha_html_report(result: HAComparison) -> bytes:
    parts = [
        "<!doctype html><html><head><meta charset='utf-8'><title>FortiGate HA Consistency</title>",
        "<style>body{background:#0b1220;color:#dce7f5;font:14px Inter,Arial,sans-serif;margin:24px}",
        "table{border-collapse:collapse;width:100%;margin:12px 0 28px}th,td{border:1px solid #263447;padding:8px;text-align:left}",
        "th{background:#172336;color:#00d4ff}.missing-in-a{background:#17432e}.missing-in-b{background:#48262b}",
        ".modified{background:#504319}a{color:#00d4ff}code{white-space:pre-wrap}",
        ".diff_add{background:#16442f;color:#bdffdc}.diff_sub{background:#48262b;color:#ffd2d2}",
        ".diff_chg{background:#504319;color:#fff0b8}.diff_header{background:#121c2b;color:#8c9bb0}",
        "</style></head><body>",
        "<h1>FortiGate HA Configuration Consistency</h1>",
        f"<p>Firewall A: {html.escape(result.before_name)}<br>Firewall B: {html.escape(result.after_name)}</p>",
        f"<h2>Health Score: {result.health_score}%</h2>",
        "<table><tr><th>Exact Matches</th><th>Equivalent Matches</th><th>Modified Matches</th><th>Missing Policies</th></tr>",
        f"<tr><td>{result.exact_policy_count}</td><td>{result.equivalent_policy_count}</td><td>{result.modified_policy_count}</td><td>{result.missing_policy_count}</td></tr></table>",
        "<h2>Object Differences</h2><table><tr><th>Category</th><th>Object</th><th>Status</th><th>Risk</th></tr>",
    ]
    for change in result.changes:
        if change.status == "Matching":
            continue
        row_class = change.status.lower().replace(" ", "-")
        parts.append(
            f'<tr class="{row_class}"><td>{html.escape(change.category)}</td>'
            f"<td>{html.escape(change.label)}</td><td>{html.escape(change.status)}</td>"
            f"<td>{html.escape(change.risk)}</td></tr>"
        )
    parts.append("</table>")
    parts.append("<h2>Finding Details</h2>")
    for change in result.changes:
        if change.status == "Matching":
            continue
        _append_html_finding(parts, change)
        parts.append(
            difflib.HtmlDiff(wrapcolumn=120).make_table(
                list(change.before_lines),
                list(change.after_lines),
                fromdesc=html.escape(f"Firewall A · {change.label}"),
                todesc=html.escape(f"Firewall B · {change.label}"),
                context=True,
                numlines=3,
            )
        )
    parts.append("<h2>Cosmetic Differences</h2>")
    for change in result.cosmetic_changes:
        _append_html_finding(parts, change)
    parts.append("<h2>Dependency Validation</h2><ul>")
    parts.extend(
        f"<li>{html.escape(item['Firewall'])}: {html.escape(item['Policy'])} references missing "
        f"{html.escape(item['Reference Type'].lower())} {html.escape(item['Missing Reference'])}</li>"
        for item in result.dependencies
    )
    parts.append("</ul><h2>Risk Analysis</h2><table><tr><th>Risk</th><th>Object</th><th>Finding</th></tr>")
    parts.extend(
        f"<tr><td>{html.escape(item['Risk'])}</td><td>{html.escape(item['Object'])}</td>"
        f"<td>{html.escape(item['Finding'])}</td></tr>"
        for item in result.risks
    )
    parts.append("</table><h2>Recommendations</h2><ul>")
    recommendations = list(
        dict.fromkeys(
            change.recommendation
            for change in result.changes
            if change.status != "Matching" and change.recommendation
        )
    )
    parts.extend(f"<li>{html.escape(item)}</li>" for item in dict.fromkeys(recommendations))
    parts.append("</ul></body></html>")
    return "".join(parts).encode("utf-8")


def build_ha_excel_report(result: HAComparison) -> BytesIO:
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Configuration Health Score", f"{result.health_score}%"])
    summary.append(["Matching Objects", result.matching_count])
    summary.append(["Different Objects", result.different_count])
    summary.append(["Missing Objects", result.missing_count])
    summary.append(["Exact Policy Matches", result.exact_policy_count])
    summary.append(["Equivalent Policy Matches", result.equivalent_policy_count])
    summary.append(["Modified Policy Matches", result.modified_policy_count])
    summary.append(["Missing Policies", result.missing_policy_count])
    summary.append(["Cosmetic Differences", result.cosmetic_count])

    changes_sheet = workbook.create_sheet("Object Differences")
    changes_sheet.append(
        [
            "Object Type",
            "Object ID",
            "Object Name",
            "Status",
            "Exact Difference",
            "Impact Assessment",
            "Risk Level",
            "Recommended Action",
            "Firewall A Configuration",
            "Firewall B Configuration",
        ]
    )
    for change in result.changes:
        changes_sheet.append(
            [
                change.category,
                change.object_id,
                change.label,
                change.status,
                "; ".join(item["Difference"] for item in exact_differences(change)),
                change.impact or change.risk_reason,
                change.risk,
                change.recommendation,
                "\n".join(change.before_lines),
                "\n".join(change.after_lines),
            ]
        )

    cosmetic_sheet = workbook.create_sheet("Cosmetic Differences")
    cosmetic_sheet.append(["Object Type", "Object Name", "Exact Difference", "Firewall A", "Firewall B"])
    for change in result.cosmetic_changes:
        cosmetic_sheet.append(
            [
                change.category,
                change.label,
                "; ".join(item["Difference"] for item in exact_differences(change)),
                "\n".join(change.before_lines),
                "\n".join(change.after_lines),
            ]
        )

    dependency_sheet = workbook.create_sheet("Dependencies")
    dependency_sheet.append(["Firewall", "Policy", "Missing Reference", "Reference Type", "Severity"])
    for item in result.dependencies:
        dependency_sheet.append([item[key] for key in ("Firewall", "Policy", "Missing Reference", "Reference Type", "Severity")])

    risk_sheet = workbook.create_sheet("Risk Analysis")
    risk_sheet.append(["Risk", "Object", "Finding"])
    for item in result.risks:
        risk_sheet.append([item[key] for key in ("Risk", "Object", "Finding")])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_ha_pdf_report(result: HAComparison) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, TableStyle

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    story = [
        Paragraph("FortiGate HA Configuration Consistency", styles["Title"]),
        Paragraph(f"Health score: {result.health_score}%", styles["Heading2"]),
        Paragraph(
            f"Functional findings: {result.different_count + result.missing_count} &nbsp; "
            f"Cosmetic differences: {result.cosmetic_count}",
            styles["BodyText"],
        ),
    ]
    rows = [["Object Type / Name", "Status / Risk", "Exact Difference", "Impact Assessment", "Recommended Action"]]
    for change in result.changes:
        if change.status == "Matching":
            continue
        exact_text = "<br/>".join(
            html.escape(
                f"{item['Setting']}: A={item['Firewall A']}; B={item['Firewall B']}"
            )
            for item in exact_differences(change)
        )
        rows.append(
            [
                Paragraph(html.escape(f"{change.category}<br/>{change.label}"), styles["BodyText"]),
                Paragraph(html.escape(f"{change.status}<br/>{change.risk}"), styles["BodyText"]),
                Paragraph(exact_text, styles["Code"]),
                Paragraph(html.escape(change.impact or change.risk_reason), styles["BodyText"]),
                Paragraph(html.escape(change.recommendation), styles["BodyText"]),
            ]
        )
    table = LongTable(rows, repeatRows=1, colWidths=[95, 70, 135, 145, 155])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#172336")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#00D4FF")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#9AAAC0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(table)
    if result.cosmetic_changes:
        story.append(Paragraph("Cosmetic Differences (excluded from score)", styles["Heading2"]))
        for change in result.cosmetic_changes:
            exact_text = "; ".join(
                f"{item['Setting']}: A={item['Firewall A']}; B={item['Firewall B']}"
                for item in exact_differences(change)
            )
            story.append(
                Paragraph(
                    html.escape(f"{change.category} {change.label}: {exact_text}"),
                    styles["BodyText"],
                )
            )
    document.build(story)
    return output.getvalue()