import unittest

from openpyxl import load_workbook

from session_analyzer import (
    analyze_sessions,
    build_session_excel_report,
    format_duration,
    parse_sessions,
)


SESSION = """session info: proto=6 proto_state=01 duration=7852 expire=30 timeout=3600 policy_id=42 state=log may_dirty serial=abc123 npu_state=0x00000001
hook=pre dir=org act=noop 172.20.109.142:50801->10.16.66.132:80
statistic(bytes/packets/allow_err): org=1048576/100/0 reply=2097152/200/0
tx speed(Bps/kbps): 1024/8 rx speed(Bps/kbps): 2048/16
"""


class SessionAnalyzerTests(unittest.TestCase):
    def test_extracts_fields_calculations_and_application(self):
        sessions = parse_sessions(SESSION)

        self.assertEqual(len(sessions), 1)
        row = sessions.iloc[0]
        self.assertEqual(row["Protocol"], "6")
        self.assertEqual(row["Proto State"], "01")
        self.assertEqual(row["Source IP"], "172.20.109.142")
        self.assertEqual(row["Source Port"], 50801)
        self.assertEqual(row["Destination IP"], "10.16.66.132")
        self.assertEqual(row["Destination Port"], 80)
        self.assertEqual(row["Application"], "HTTP")
        self.assertEqual(row["Policy ID"], "42")
        self.assertEqual(row["Serial Number"], "abc123")
        self.assertEqual(row["Duration Human Format"], "2h 10m 52s")
        self.assertEqual(row["Total Bytes"], 3 * 1024 * 1024)
        self.assertEqual(row["Upload Bytes"], 1024 * 1024)
        self.assertEqual(row["Download Bytes"], 2 * 1024 * 1024)
        self.assertEqual(row["Current Throughput"], 24)
        self.assertEqual(row["Offload Status"], "Offloaded")

    def test_builds_health_findings_and_incident_summary(self):
        input_text = SESSION.replace("duration=7852", "duration=14401").replace(
            "org=1048576/100/0 reply=2097152/200/0",
            "org=1200000000/100/0 reply=100/200/0",
        ).replace("npu_state=0x00000001", "no_ofld_reason: disabled")
        analysis = analyze_sessions(input_text)

        self.assertEqual(analysis.summary["Total Sessions"], 1)
        self.assertEqual(analysis.summary["HTTP Sessions"], 1)
        self.assertEqual(analysis.summary["Sessions Without NPU Offload"], 1)
        self.assertEqual(
            set(analysis.health_findings["Finding"]),
            {"LONG RUNNING", "VERY HIGH BANDWIDTH", "NOT OFFLOADED"},
        )
        self.assertIn("Recommendations:", analysis.incident_summary)

    def test_parses_two_hundred_sessions_and_exports_all_sheets(self):
        text = "".join(
            SESSION.replace("serial=abc123", f"serial=session-{index}")
            for index in range(200)
        )
        analysis = analyze_sessions(text)

        self.assertEqual(len(analysis.sessions), 200)
        self.assertEqual(len(analysis.top_talkers), 20)
        workbook = load_workbook(build_session_excel_report(analysis), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Summary",
                "All Sessions",
                "Top Talkers",
                "Longest Sessions",
                "Source Analysis",
                "Destination Analysis",
                "Health Findings",
            ],
        )

    def test_empty_input_and_duration_format(self):
        self.assertTrue(parse_sessions("").empty)
        self.assertEqual(format_duration(0), "0s")
        self.assertEqual(format_duration(7852), "2h 10m 52s")


if __name__ == "__main__":
    unittest.main()