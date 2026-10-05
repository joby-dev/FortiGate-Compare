"""Line-oriented FortiGate configuration diff helpers."""

from __future__ import annotations

import difflib
import html
import json
import re
from bisect import bisect_right
from dataclasses import dataclass
from io import BytesIO
from openpyxl import Workbook


SECTION_FILTERS = (
    "Firewall Policies",
    "Proxy Policies",
    "Local-In Policies",
    "Address Objects",
    "Service Objects",
    "VIP Objects",
    "Static Routes",
    "SD-WAN",
    "SSL Profiles",
    "Web Filter Profiles",
)

SECTION_SUMMARIES = {
    "Firewall Policies": "Firewall Policies Changed",
    "Proxy Policies": "Proxy Policies Changed",
    "Local-In Policies": "Local-In Policies Changed",
    "IPv6 Firewall Policies": "IPv6 Firewall Policies Changed",
    "IPv6 Proxy Policies": "IPv6 Proxy Policies Changed",
    "Address Objects": "Address Objects Changed",
    "Service Objects": "Services Changed",
    "VIP Objects": "VIP Objects Changed",
    "Static Routes": "Routes Changed",
    "SD-WAN": "SD-WAN Changed",
    "SSL Profiles": "SSL Profiles Changed",
    "Web Filter Profiles": "Web Filter Profiles Changed",
}

SECTION_LABELS = {
    "Firewall Policies": "Firewall Policy",
    "Proxy Policies": "Proxy Policy",
    "Local-In Policies": "Local-In Policy",
    "IPv6 Firewall Policies": "IPv6 Firewall Policy",
    "IPv6 Proxy Policies": "IPv6 Proxy Policy",
    "Address Objects": "Address Object",
    "Service Objects": "Service Object",
    "VIP Objects": "VIP Object",
    "Static Routes": "Static Route",
    "SD-WAN": "SD-WAN",
    "SSL Profiles": "SSL Profile",
    "Web Filter Profiles": "Web Filter Profile",
}
POLICY_TYPES_BY_SECTION = {
    "firewall policy": "Firewall",
    "firewall proxy-policy": "Proxy",
    "firewall local-in-policy": "Local-In",
    "firewall policy6": "IPv6 Firewall",
    "firewall proxy-policy6": "IPv6 Proxy",
}


@dataclass(frozen=True)
class ObjectSpan:
    start: int
    end: int
    section: str
    category: str | None
    object_id: str
    name: str

    @property
    def key(self) -> tuple[str, str]:
        return (self.section, self.object_id)

    @property
    def label(self) -> str:
        if self.category is None:
            return f"{self.section}: {self.name}"
        return f"{SECTION_LABELS[self.category]} {self.object_id}"


@dataclass(frozen=True)
class DiffRow:
    kind: str
    before_line: int | None
    after_line: int | None
    before_text: str
    after_text: str


@dataclass
class DiffHunk:
    title: str
    categories: set[str]
    objects: set[tuple[str, str]]
    rows: list[DiffRow]


@dataclass
class RawDiffResult:
    before_name: str
    after_name: str
    before_lines: list[str]
    after_lines: list[str]
    hunks: list[DiffHunk]
    changed_objects: list[dict[str, str]]
    summary: dict[str, int]
    order_changes: list[str]
    object_statuses: dict[tuple[str, str], str]
    object_summary: dict[str, int]
    full_rows: list[DiffRow]
    before_spans: list[ObjectSpan]
    after_spans: list[ObjectSpan]


def _section_category(section: str) -> str | None:
    normalized = section.lower()
    if normalized == "firewall policy":
        return "Firewall Policies"
    if normalized == "firewall proxy-policy":
        return "Proxy Policies"
    if normalized == "firewall local-in-policy":
        return "Local-In Policies"
    if normalized == "firewall policy6":
        return "Firewall Policies"
    if normalized == "firewall proxy-policy6":
        return "Proxy Policies"
    if normalized in {"firewall address", "firewall addrgrp"}:
        return "Address Objects"
    if normalized.startswith("firewall service "):
        return "Service Objects"
    if normalized in {"firewall vip", "firewall vipgrp"}:
        return "VIP Objects"
    if normalized == "router static":
        return "Static Routes"
    if normalized in {"system sdwan", "system virtual-wan-link"}:
        return "SD-WAN"
    if normalized == "firewall ssl-ssh-profile":
        return "SSL Profiles"
    if normalized == "webfilter profile":
        return "Web Filter Profiles"
    return None


def _is_non_functional_line(line: str) -> bool:
    normalized = line.strip()
    if not normalized:
        return False
    if normalized.startswith(("#", "!", "//")):
        return True
    return bool(
        re.match(
            r"^(?:set\s+)?(?:config-version|build(?:no|number)?|conf_file_ver|"
            r"revision(?:-number)?|timestamp|last-modified|admin[\s_-]+session|session-id)\b",
            normalized,
            flags=re.IGNORECASE,
        )
        or re.match(r"^set\s+comments?\b", normalized, flags=re.IGNORECASE)
    )


def _mask_ignored_lines(lines: list[str], ignore_categories: set[str]) -> list[str]:
    config_stack: list[str] = []
    masked_lines = []
    for line in lines:
        stripped = line.strip()
        normalized = stripped.casefold()
        config_match = re.match(r"^config\s+(.+)$", stripped, flags=re.IGNORECASE)
        if config_match:
            config_stack.append(re.sub(r"\s+", " ", config_match.group(1)).casefold())

        in_interface = "system interface" in config_stack
        in_ha = "system ha" in config_stack
        ignored = (
            "Hostname" in ignore_categories
            and re.match(r"^set\s+hostname\b", normalized) is not None
        ) or (
            "Interface IP" in ignore_categories
            and in_interface
            and re.match(r"^set\s+ip6?\b", normalized) is not None
        ) or (
            "UUID" in ignore_categories
            and re.match(r"^set\s+uuid(?:-index)?\b", normalized) is not None
        ) or (
            "Comments" in ignore_categories
            and (
                normalized.startswith(("#", "!", "//"))
                or re.match(r"^set\s+comments?\b", normalized) is not None
            )
        ) or (
            "Config Revisions" in ignore_categories
            and _is_non_functional_line(line)
        ) or ("HA Settings" in ignore_categories and in_ha)

        masked_lines.append("\n" if ignored else line)
        if normalized == "end" and config_stack:
            config_stack.pop()
    return masked_lines


def _object_spans(lines: list[str]) -> list[ObjectSpan]:
    spans: list[ObjectSpan] = []
    section: str | None = None
    section_depth = 0
    current_start: int | None = None
    current_id = ""
    current_name = ""

    def finish_object(end: int) -> None:
        nonlocal current_start, current_id, current_name
        if section is not None and current_start is not None:
            spans.append(
                ObjectSpan(
                    current_start,
                    max(current_start, end),
                    section,
                    _section_category(section),
                    current_id,
                    current_name or current_id,
                )
            )
        current_start = None
        current_id = ""
        current_name = ""

    for index, raw_line in enumerate(lines):
        line = raw_line.strip()
        config_match = re.match(r"^config\s+(.+)$", line, flags=re.IGNORECASE)
        if section is None and config_match:
            section = config_match.group(1).strip()
            section_depth = 1
            continue
        if section is None:
            continue
        if config_match:
            section_depth += 1
            continue

        edit_match = re.match(r"^edit\s+(.+)$", line, flags=re.IGNORECASE)
        if edit_match:
            finish_object(index - 1)
            current_start = index
            current_id = edit_match.group(1).strip().strip('"')
            current_name = current_id
            continue

        name_match = re.match(r'^set\s+name\s+"?(.*?)"?$', line, flags=re.IGNORECASE)
        if name_match and current_start is not None:
            current_name = name_match.group(1).strip()
            continue

        if line.lower() == "next":
            finish_object(index)
            continue
        if line.lower() == "end":
            finish_object(index - 1)
            section_depth -= 1
            if section_depth <= 0:
                section = None

    if section is not None:
        finish_object(len(lines) - 1)
    return spans


def _spans_for_range(
    spans: list[ObjectSpan], starts: list[int], start: int, end: int
) -> list[ObjectSpan]:
    if start == end:
        index = bisect_right(starts, max(0, start - 1)) - 1
        return [spans[index]] if index >= 0 and spans[index].end >= start else []
    index = max(0, bisect_right(starts, start) - 1)
    matches = []
    while index < len(spans) and spans[index].start < end:
        if spans[index].end >= start:
            matches.append(spans[index])
        index += 1
    return matches


def _build_rows(
    tag: str,
    before_lines: list[str],
    after_lines: list[str],
    before_start: int,
    before_end: int,
    after_start: int,
    after_end: int,
) -> list[DiffRow]:
    before_slice = before_lines[before_start:before_end]
    after_slice = after_lines[after_start:after_end]
    rows: list[DiffRow] = []

    if tag == "equal":
        rows.extend(
            DiffRow(
                "Unchanged",
                before_start + offset + 1,
                after_start + offset + 1,
                before_line,
                after_line,
            )
            for offset, (before_line, after_line) in enumerate(
                zip(before_slice, after_slice)
            )
        )
    elif tag == "replace":
        paired = min(len(before_slice), len(after_slice))
        rows.extend(
            DiffRow(
                "Modified",
                before_start + offset + 1,
                after_start + offset + 1,
                before_slice[offset],
                after_slice[offset],
            )
            for offset in range(paired)
        )
        rows.extend(
            DiffRow(
                "Removed",
                before_start + offset + 1,
                None,
                before_slice[offset],
                "",
            )
            for offset in range(paired, len(before_slice))
        )
        rows.extend(
            DiffRow(
                "Added",
                None,
                after_start + offset + 1,
                "",
                after_slice[offset],
            )
            for offset in range(paired, len(after_slice))
        )
    elif tag == "delete":
        rows.extend(
            DiffRow("Removed", before_start + offset + 1, None, line, "")
            for offset, line in enumerate(before_slice)
        )
    elif tag == "insert":
        rows.extend(
            DiffRow("Added", None, after_start + offset + 1, "", line)
            for offset, line in enumerate(after_slice)
        )
    return rows


def _policy_order_changes(
    before_spans: list[ObjectSpan], after_spans: list[ObjectSpan]
) -> tuple[list[str], dict[str, set[str]]]:
    changes = []
    moved_by_section: dict[str, set[str]] = {}
    for section, title in (
        ("firewall policy", "Firewall Policy"),
        ("firewall proxy-policy", "Proxy Policy"),
        ("firewall local-in-policy", "Local-In Policy"),
        ("firewall policy6", "IPv6 Firewall Policy"),
        ("firewall proxy-policy6", "IPv6 Proxy Policy"),
    ):
        before_order = [span.object_id for span in before_spans if span.section.lower() == section]
        after_order = [span.object_id for span in after_spans if span.section.lower() == section]
        common = set(before_order) & set(after_order)
        before_common = [policy_id for policy_id in before_order if policy_id in common]
        after_common = [policy_id for policy_id in after_order if policy_id in common]
        if before_common != after_common:
            changes.append(f"{title} order changed")
            moved_by_section[section] = {
                before_common[index]
                for index in range(len(before_common))
                if before_common[index] != after_common[index]
            } | {
                after_common[index]
                for index in range(len(after_common))
                if before_common[index] != after_common[index]
            }
    return changes, moved_by_section


def compare_raw_configurations(
    before_text: str,
    after_text: str,
    before_name: str = "Pre-Change Configuration",
    after_name: str = "Post-Change Configuration",
    ignore_non_functional_changes: bool = False,
    ignore_categories: set[str] | None = None,
) -> RawDiffResult:
    """Compare complete configs line-by-line and retain changed hunks."""
    before_lines = before_text.splitlines(keepends=True)
    after_lines = after_text.splitlines(keepends=True)
    selected_ignores = set(ignore_categories or ())
    if ignore_non_functional_changes:
        selected_ignores.update({"Comments", "Config Revisions"})
    if selected_ignores:
        before_lines = _mask_ignored_lines(before_lines, selected_ignores)
        after_lines = _mask_ignored_lines(after_lines, selected_ignores)
    matcher = difflib.SequenceMatcher(
        None, before_lines, after_lines, autojunk=True
    )
    all_opcodes = matcher.get_opcodes()
    full_rows = [
        row
        for tag, before_start, before_end, after_start, after_end in all_opcodes
        for row in _build_rows(
            tag,
            before_lines,
            after_lines,
            before_start,
            before_end,
            after_start,
            after_end,
        )
    ]
    before_spans = _object_spans(before_lines)
    after_spans = _object_spans(after_lines)
    before_starts = [span.start for span in before_spans]
    after_starts = [span.start for span in after_spans]
    grouped_opcodes = matcher.get_grouped_opcodes(3)
    hunks: list[DiffHunk] = []
    object_records: dict[tuple[str, str], dict[str, str]] = {}
    changed_counts: dict[str, set[tuple[str, str]]] = {}

    for grouped in grouped_opcodes:
        rows: list[DiffRow] = []
        objects: set[tuple[str, str, str]] = set()
        categories: set[str] = set()
        for tag, i1, i2, j1, j2 in grouped:
            rows.extend(_build_rows(tag, before_lines, after_lines, i1, i2, j1, j2))
            if tag == "equal":
                continue
            changed_spans = _spans_for_range(
                before_spans, before_starts, i1, i2
            ) + _spans_for_range(after_spans, after_starts, j1, j2)
            if not changed_spans:
                changed_spans = _spans_for_range(
                    before_spans, before_starts, i1, i1
                ) + _spans_for_range(after_spans, after_starts, j1, j1)
            for span in changed_spans:
                if span.category is None:
                    continue
                objects.add(span.key)
                categories.add(span.category)
                object_records[span.key] = {
                    "Object": span.label,
                    "Section": span.section,
                    "Category": span.category,
                    "Policy Type": POLICY_TYPES_BY_SECTION.get(
                        span.section.lower(), span.category
                    ),
                    "Object ID": span.object_id,
                    "Object Name": span.name,
                }
                changed_counts.setdefault(span.category, set()).add(span.key)

        object_title = ", ".join(sorted(object_records[key]["Object"] for key in objects))
        title = object_title or f"Changed lines {grouped[0][1] + 1}-{grouped[-1][2]}"
        hunks.append(DiffHunk(title, categories, objects, rows))

    summary = {label: 0 for label in SECTION_SUMMARIES.values()}
    for category, keys in changed_counts.items():
        summary[SECTION_SUMMARIES.get(category, f"{category} Changed")] = len(keys)
    order_changes, moved_by_section = _policy_order_changes(
        before_spans, after_spans
    )
    for section, moved_ids in moved_by_section.items():
        spans = [
            span
            for span in before_spans + after_spans
            if span.section.lower() == section and span.object_id in moved_ids
        ]
        for span in spans:
            if span.category is None:
                continue
            object_records[span.key] = {
                "Object": span.label,
                "Section": span.section,
                "Category": span.category,
                "Policy Type": POLICY_TYPES_BY_SECTION.get(
                    span.section.lower(), span.category
                ),
                "Object ID": span.object_id,
                "Object Name": span.name,
            }
            changed_counts.setdefault(span.category, set()).add(span.key)
    moved_spans = {
        span.key: span
        for span in before_spans + after_spans
        if span.object_id in moved_by_section.get(span.section.lower(), set())
    }
    for span in moved_spans.values():
        if span.category is None or not hunks:
            continue
        compatible_hunks = [hunk for hunk in hunks if span.category in hunk.categories]
        if not compatible_hunks:
            continue
        span_line = span.start + 1
        nearest_hunk = min(
            compatible_hunks,
            key=lambda hunk: min(
                (
                    abs(line_number - span_line)
                    for row in hunk.rows
                    for line_number in (row.before_line, row.after_line)
                    if line_number is not None
                ),
                default=0,
            ),
        )
        nearest_hunk.objects.add(span.key)
        nearest_hunk.title = ", ".join(
            sorted(
                object_records[key]["Object"]
                for key in nearest_hunk.objects
                if key in object_records
            )
        )
    for category, keys in changed_counts.items():
        summary[SECTION_SUMMARIES.get(category, f"{category} Changed")] = len(keys)
    changed_objects = sorted(
        object_records.values(), key=lambda item: (item["Category"], item["Object ID"])
    )
    before_object_keys = {span.key for span in before_spans if span.category is not None}
    after_object_keys = {span.key for span in after_spans if span.category is not None}
    object_statuses = {}
    for item in changed_objects:
        key = (item["Section"], item["Object ID"])
        if key not in before_object_keys:
            object_statuses[key] = "Missing in Pre-Change"
        elif key not in after_object_keys:
            object_statuses[key] = "Missing in Post-Change"
        else:
            object_statuses[key] = "Modified"
    object_summary = {
        "Matching Objects": len((before_object_keys | after_object_keys) - set(object_statuses)),
        "Modified Objects": sum(status == "Modified" for status in object_statuses.values()),
        "Missing Objects": sum(status.startswith("Missing in") for status in object_statuses.values()),
    }

    return RawDiffResult(
        before_name,
        after_name,
        before_lines,
        after_lines,
        hunks,
        changed_objects,
        summary,
        order_changes,
        object_statuses,
        object_summary,
        full_rows,
        before_spans,
        after_spans,
    )


def render_side_by_side_html(hunk: DiffHunk, max_rows: int = 1000) -> str:
    """Return an escaped two-column HTML table for one hunk."""
    colors = {
        "Added": "#D4F8D4",
        "Removed": "#FFD6D6",
        "Modified": "#FFF1B8",
        "Unchanged": "#FFFFFF",
    }
    rows = hunk.rows[:max_rows]
    output = [
        '<table class="diff"><thead><tr><th>Line</th><th>Pre-Change</th>'
        "<th>Line</th><th>Post-Change</th></tr></thead><tbody>"
    ]
    for row in rows:
        color = colors[row.kind]
        before_text = html.escape(row.before_text.rstrip())
        after_text = html.escape(row.after_text.rstrip())
        if row.kind == "Modified":
            before_text, after_text = _inline_change_parts(
                row.before_text.rstrip(), row.after_text.rstrip()
            )
        output.append(
            f'<tr class="{row.kind.lower()}" style="background:{color}">'
            f'<td>{row.before_line or ""}</td>'
            f"<td><pre>{before_text}</pre></td>"
            f'<td>{row.after_line or ""}</td>'
            f"<td><pre>{after_text}</pre></td></tr>"
        )
    output.append("</tbody></table>")
    if len(hunk.rows) > max_rows:
        output.append(
            f"<p>Showing {max_rows:,} of {len(hunk.rows):,} hunk lines. "
            "The full change is included in the exported report.</p>"
        )
    return "".join(output)


def build_raw_html_report(result: RawDiffResult) -> bytes:
    """Create an HTML report of changed hunks and policy order changes."""
    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        "<title>FortiGate Configuration Diff</title>",
        "<style>body{font-family:Segoe UI,Arial,sans-serif;margin:2rem;"
        "background:#EAF6FF;color:#20252a}a{color:#1769aa}"
        "table.diff{border-collapse:collapse;width:100%;margin:1rem 0;table-layout:fixed}"
        "td,th{border:1px solid #B8DFF5;padding:.35rem;text-align:left;vertical-align:top}"
        "th{position:sticky;top:0;background:#D6EEFF}"
        "td:nth-child(1),td:nth-child(3){width:4rem;color:#59636d;text-align:right}"
        "pre{white-space:pre;overflow-x:auto;margin:0;"
        "font-family:Consolas,'Cascadia Code',monospace;color:#20252a}"
        ".inline-added,.inline-removed{background:#FFF1B8;color:#20252a}"
        ".added{background:#D4F8D4}.removed{background:#FFD6D6}.modified{background:#FFF1B8}"
        "</style></head><body id='top'>",
        "<h1>FortiGate Configuration Diff</h1>",
        f"<p>Before: {html.escape(result.before_name)}<br>After: "
        f"{html.escape(result.after_name)}</p>",
        "<nav><h2>Changed Sections</h2><ul>",
    ]
    parts.extend(
        f'<li><a href="#change-{index}">{html.escape(hunk.title)}</a></li>'
        for index, hunk in enumerate(result.hunks, start=1)
    )
    parts.extend(
        [
            "</ul></nav>",
        "<h2>Configuration Changes Summary</h2><ul>",
        ]
    )
    parts.extend(
        f"<li>{html.escape(label)}: {count:,}</li>"
        for label, count in result.summary.items()
    )
    parts.append("</ul>")
    if result.order_changes:
        parts.append("<h2>Order Changes</h2><ul>")
        parts.extend(f"<li>{html.escape(change)}</li>" for change in result.order_changes)
        parts.append("</ul>")
    for index, hunk in enumerate(result.hunks, start=1):
        parts.append(
            f'<h2 id="change-{index}">{html.escape(hunk.title)} '
            f'<a href="#top">Back to top</a></h2>'
        )
        parts.append(render_side_by_side_html(hunk, max_rows=len(hunk.rows)))
    parts.append("</body></html>")
    return "".join(parts).encode("utf-8")


def build_raw_excel_report(result: RawDiffResult) -> BytesIO:
    """Create a write-only Excel report with summary, objects, and changed lines."""
    workbook = Workbook(write_only=True)
    summary_sheet = workbook.create_sheet("Summary")
    summary_sheet.append(["Metric", "Count"])
    for label, count in result.summary.items():
        summary_sheet.append([label, count])
    total_changed_lines = sum(
        row.kind != "Unchanged" for hunk in result.hunks for row in hunk.rows
    )
    summary_sheet.append(["Total Changed Lines", total_changed_lines])
    summary_sheet.append(["Policy Order Changed", bool(result.order_changes)])
    for order_change in result.order_changes:
        summary_sheet.append(["Order Change", order_change])

    object_sheet = workbook.create_sheet("Changed Objects")
    object_sheet.append(
        ["Object", "Section", "Category", "Policy Type", "Object ID", "Object Name"]
    )
    for item in result.changed_objects:
        object_sheet.append(
            [
                item[key]
                for key in (
                    "Object",
                    "Section",
                    "Category",
                    "Policy Type",
                    "Object ID",
                    "Object Name",
                )
            ]
        )

    diff_sheet = workbook.create_sheet("Line Diff")
    diff_sheet.append(
        ["Changed Object", "Change", "Line Number Before", "Line Number After", "Before", "After"]
    )
    written_rows = 0
    max_rows = 1_048_575
    for hunk in result.hunks:
        for row in hunk.rows:
            if row.kind == "Unchanged":
                continue
            if written_rows >= max_rows:
                break
            diff_sheet.append(
                [
                    hunk.title,
                    row.kind,
                    row.before_line,
                    row.after_line,
                    row.before_text.rstrip("\r\n"),
                    row.after_text.rstrip("\r\n"),
                ]
            )
            written_rows += 1
        if written_rows >= max_rows:
            break
    if total_changed_lines > written_rows:
        summary_sheet.append(["Excel Diff Rows Omitted", total_changed_lines - written_rows])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_raw_pdf_report(result: RawDiffResult) -> bytes:
    """Create a landscape PDF with a summary and color-coded changed lines."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, TableStyle

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=28,
    )
    styles = getSampleStyleSheet()
    code_style = ParagraphStyle(
        "DiffCode",
        parent=styles["BodyText"],
        fontName="Courier",
        fontSize=7,
        leading=8,
        textColor=colors.HexColor("#202124"),
        wordWrap="CJK",
    )
    cell_style = ParagraphStyle(
        "DiffCell",
        parent=code_style,
        fontName="Helvetica",
        fontSize=7,
        leading=8,
    )
    story = [
        Paragraph("FortiGate Configuration Diff", styles["Title"]),
        Paragraph(
            f"Pre-Change: {html.escape(result.before_name)}<br/>"
            f"Post-Change: {html.escape(result.after_name)}",
            styles["BodyText"],
        ),
        Paragraph("Change Summary", styles["Heading2"]),
    ]
    summary_rows = [["Section", "Changed Objects"]]
    summary_rows.extend(
        [html.escape(label), f"{count:,}"]
        for label, count in result.summary.items()
        if count
    )
    story.append(
        LongTable(
            summary_rows,
            colWidths=[240, 100],
            repeatRows=1,
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#25282d")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aab0b8")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ]
            ),
        )
    )
    story.append(Paragraph("Changed Lines", styles["Heading2"]))
    diff_rows = [
        [
            "Section / Object",
            "Change",
            "Before",
            "After",
            "Pre-Change",
            "Post-Change",
        ]
    ]
    row_kinds = []
    for hunk in result.hunks:
        for row in hunk.rows:
            if row.kind == "Unchanged":
                continue
            diff_rows.append(
                [
                    Paragraph(html.escape(hunk.title), cell_style),
                    row.kind,
                    str(row.before_line or ""),
                    str(row.after_line or ""),
                    Paragraph(html.escape(row.before_text.rstrip()), code_style),
                    Paragraph(html.escape(row.after_text.rstrip()), code_style),
                ]
            )
            row_kinds.append(row.kind)
    diff_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#25282d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aab0b8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    row_colors = {
        "Added": colors.HexColor("#d8f0df"),
        "Removed": colors.HexColor("#f3d8d8"),
        "Modified": colors.HexColor("#f6edc7"),
    }
    for row_index, kind in enumerate(row_kinds, start=1):
        diff_style.append(
            ("BACKGROUND", (0, row_index), (-1, row_index), row_colors[kind])
        )
    story.append(
        LongTable(
            diff_rows,
            colWidths=[142, 48, 42, 42, 255, 255],
            repeatRows=1,
            style=TableStyle(diff_style),
        )
    )
    document.build(story)
    return output.getvalue()


def build_unified_diff(result: RawDiffResult, max_lines: int | None = None) -> str:
    """Return a unified diff, optionally limiting the interactive preview."""
    lines = difflib.unified_diff(
        result.before_lines,
        result.after_lines,
        fromfile=result.before_name,
        tofile=result.after_name,
    )
    if max_lines is None:
        return "".join(lines)
    output = []
    for line_number, line in enumerate(lines):
        if line_number >= max_lines:
            output.append("... unified diff preview truncated; full diff is in the exports ...\n")
            break
        output.append(line)
    return "".join(output)


def _inline_change_parts(before_text: str, after_text: str) -> tuple[str, str]:
    before_parts = []
    after_parts = []
    matcher = difflib.SequenceMatcher(None, before_text, after_text, autojunk=False)
    for tag, before_start, before_end, after_start, after_end in matcher.get_opcodes():
        before_part = html.escape(before_text[before_start:before_end])
        after_part = html.escape(after_text[after_start:after_end])
        if tag == "equal":
            before_parts.append(before_part)
            after_parts.append(after_part)
        else:
            if before_part:
                before_parts.append(f'<span class="inline-removed">{before_part}</span>')
            if after_part:
                after_parts.append(f'<span class="inline-added">{after_part}</span>')
    return "".join(before_parts), "".join(after_parts)


def render_side_by_side_document(
    hunk: DiffHunk,
    before_name: str,
    after_name: str,
    visible_kinds: set[str] | None = None,
    focus_row: int | None = None,
    max_rows: int = 1200,
    before_label: str = "Firewall A / Pre-Change",
    after_label: str = "Firewall B / Post-Change",
    source_before_lines: list[str] | None = None,
    source_after_lines: list[str] | None = None,
    storage_key: str = "fortigate-diff",
    navigation_target: dict[str, object] | None = None,
    navigation_token: str | None = None,
) -> str:
    """Build a virtualized Monaco diff viewer with SequenceMatcher navigation."""
    visible_rows = [
        row
        for row in hunk.rows
        if visible_kinds is None or row.kind in visible_kinds
    ]
    if source_before_lines is None or source_after_lines is None:
        visible_rows = visible_rows[:max_rows]
        before_numbers = [row.before_line for row in visible_rows if row.before_line is not None]
        after_numbers = [row.after_line for row in visible_rows if row.after_line is not None]
        before_lines = [
            row.before_text.rstrip("\r\n")
            for row in visible_rows
            if row.before_line is not None
        ]
        after_lines = [
            row.after_text.rstrip("\r\n")
            for row in visible_rows
            if row.after_line is not None
        ]
    else:
        before_lines = [line.rstrip("\r\n") for line in source_before_lines]
        after_lines = [line.rstrip("\r\n") for line in source_after_lines]
        before_numbers = list(range(1, len(before_lines) + 1))
        after_numbers = list(range(1, len(after_lines) + 1))

    matcher = difflib.SequenceMatcher(
        None, before_lines, after_lines, autojunk=True
    )
    difference_blocks = []
    line_map = []
    for tag, a_start, a_end, b_start, b_end in matcher.get_opcodes():
        a_original = before_numbers[a_start] if a_start < len(before_numbers) else (
            before_numbers[-1] + 1 if before_numbers else 1
        )
        b_original = after_numbers[b_start] if b_start < len(after_numbers) else (
            after_numbers[-1] + 1 if after_numbers else 1
        )
        line_map.append(
            {
                "aStart": a_start + 1,
                "aEnd": a_end,
                "bStart": b_start + 1,
                "bEnd": b_end,
                "aOriginalStart": a_original,
                "bOriginalStart": b_original,
                "aLength": a_end - a_start,
                "bLength": b_end - b_start,
            }
        )
        if tag != "equal":
            difference_blocks.append(
                {
                    "aStart": a_start + 1,
                    "aEnd": a_end,
                    "bStart": b_start + 1,
                    "bEnd": b_end,
                    "aOriginalStart": a_original,
                    "bOriginalStart": b_original,
                }
            )

    def to_json(value: object) -> str:
        return (
            json.dumps(value, ensure_ascii=True, separators=(",", ":"))
            .replace("<", "\\u003c")
            .replace(">", "\\u003e")
            .replace("&", "\\u0026")
        )

    before_text = "\n".join(before_lines)
    after_text = "\n".join(after_lines)
    blocks_json = to_json(difference_blocks)
    line_map_json = to_json(line_map)
    before_json = to_json(before_text)
    after_json = to_json(after_text)
    target_json = to_json(navigation_target or {})
    initial_index = min(focus_row or 0, max(0, len(difference_blocks) - 1))
    output = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width,initial-scale=1'>",
        "<style>*{box-sizing:border-box}html,body{height:100%;margin:0;overflow:hidden;",
        "background:#fff;color:#20252a;font:13px 'Segoe UI',sans-serif}.layout{height:100vh;",
        "display:grid;grid-template-columns:72px minmax(0,1fr);grid-template-rows:42px minmax(0,1fr)}",
        ".toolbar{grid-column:1/-1;display:flex;align-items:center;gap:8px;padding:5px 10px;",
        "border-bottom:1px solid #b8dff5;background:#eaf6ff}.toolbar button,.navigator button{",
        "border:1px solid #b8dff5;background:#fff;color:#20252a;border-radius:3px;padding:5px 9px;",
        "cursor:pointer}.toolbar button:hover,.navigator button:hover{border-color:#1769aa;color:#1769aa}",
        "#counter{font:12px Consolas,monospace;min-width:130px;text-align:center}.mapping{",
        "color:#59636d;font:12px Consolas,monospace;margin-left:auto}.file-label{max-width:30vw;",
        "overflow:hidden;text-overflow:ellipsis;white-space:nowrap;color:#59636d}.navigator{overflow:auto;",
        "border-right:1px solid #b8dff5;padding:6px;display:flex;flex-direction:column;gap:4px}",
        ".navigator button{font:12px Consolas,monospace;text-align:left;padding:5px 7px}",
        ".navigator button.active{background:#1769aa;color:#fff;border-color:#1769aa}#editor{",
        "width:100%;height:100%;min-width:0;min-height:0}.loading{padding:18px;color:#59636d}",
        ".selected-object-flash{background:rgba(255,210,64,.32)!important;outline:1px solid #d28b00}",
        ".missing-policy-marker{padding:3px 9px;color:#742d2d;background:#ffe1e1;border-left:3px solid #c33;font:12px Consolas,monospace}",
        "</style></head><body><div class='layout'><div class='toolbar'>",
        "<button id='previous' type='button'>Previous Difference</button>",
        "<span id='counter'>Difference 0 of 0</span>",
        "<button id='next' type='button'>Next Difference</button>",
        "<label class='sync-control'><input id='sync-scrolling' type='checkbox'> Sync Scrolling</label>",
        "<span class='mapping' id='mapping'></span><span class='file-label'>",
        html.escape(before_label.upper()), ": ", html.escape(before_name), " | ",
        html.escape(after_label.upper()), ": ", html.escape(after_name),
        "</span></div><nav class='navigator' id='navigator' aria-label='Difference navigator'></nav>",
        "<main id='editor'><div class='loading'>Loading Monaco diff editor...</div></main></div>",
        "<script src='https://cdn.jsdelivr.net/npm/monaco-editor@0.52.2/min/vs/loader.js'></script>",
        "<script>",
        f"const beforeText={before_json},afterText={after_json};",
        f"const blocks={blocks_json},lineMap={line_map_json};",
        f"const navigationTarget={target_json};",
        f"const navigationToken={to_json(navigation_token or '')};",
        f"const storageKey={to_json(storage_key)};let activeIndex=-1;let restoring=true;let syncing=false;",
        "const container=document.getElementById('editor'),navigatorPanel=document.getElementById('navigator');",
        "const counter=document.getElementById('counter'),mapping=document.getElementById('mapping');",
        "if(!window.require){container.innerHTML='<div class=\"loading\">Monaco Editor could not load. Check network access to cdn.jsdelivr.net.</div>';}",
        "else{require.config({paths:{vs:'https://cdn.jsdelivr.net/npm/monaco-editor@0.52.2/min/vs'}});",
        "window.MonacoEnvironment={getWorkerUrl:()=> 'data:text/javascript;charset=utf-8,'+encodeURIComponent('importScripts(\"https://cdn.jsdelivr.net/npm/monaco-editor@0.52.2/min/vs/base/worker/workerMain.js\");')};",
        "require(['vs/editor/editor.main'],()=>{container.innerHTML='';",
        "const diff=monaco.editor.createDiffEditor(container,{readOnly:true,originalEditable:false,",
        "renderSideBySide:true,automaticLayout:true,wordWrap:'off',minimap:{enabled:false},",
        "scrollBeyondLastLine:false,renderOverviewRuler:true,ignoreTrimWhitespace:false});",
        "const original=monaco.editor.createModel(beforeText,'ini'),modified=monaco.editor.createModel(afterText,'ini');",
        "diff.setModel({original,modified});const left=diff.getOriginalEditor(),right=diff.getModifiedEditor();",
        "const syncToggle=document.getElementById('sync-scrolling'),saved=JSON.parse(localStorage.getItem(storageKey)||'{}');",
        "syncToggle.checked=!!saved.sync;",
        "const navigatorButtons=[];blocks.forEach((block,index)=>{const button=document.createElement('button');",
        "button.type='button';button.textContent='#'+(index+1);button.title='Change block #'+(index+1);",
        "button.addEventListener('click',()=>goTo(index));navigatorPanel.appendChild(button);navigatorButtons.push(button)});",
        "function goTo(index){if(!blocks.length)return;activeIndex=(index+blocks.length)%blocks.length;",
        "const block=blocks[activeIndex],a=Math.max(1,Math.min(left.getModel().getLineCount(),block.aStart)),",
        "b=Math.max(1,Math.min(right.getModel().getLineCount(),block.bStart));",
        "left.revealLineInCenter(a);right.revealLineInCenter(b);",
        "counter.textContent='Change Block #'+(activeIndex+1)+' | Difference '+(activeIndex+1)+' of '+blocks.length;",
        "navigatorButtons.forEach((button,i)=>button.classList.toggle('active',i===activeIndex));",
        "showMapping(a,true);}",
        "function showMapping(line,isOriginal){let segment;if(isOriginal){segment=lineMap.find(item=>line>=item.aStart&&line<=Math.max(item.aStart,item.aEnd));",
        "if(segment){const offset=Math.max(0,line-segment.aStart),mapped=segment.bLength?Math.min(offset,segment.bLength-1):0;",
        "mapping.textContent='A '+(segment.aOriginalStart+offset)+' -> B '+(segment.bOriginalStart+mapped);return;}}",
        "else{segment=lineMap.find(item=>line>=item.bStart&&line<=Math.max(item.bStart,item.bEnd));",
        "if(segment){const offset=Math.max(0,line-segment.bStart),mapped=segment.aLength?Math.min(offset,segment.aLength-1):0;",
        "mapping.textContent='A '+(segment.aOriginalStart+mapped)+' -> B '+(segment.bOriginalStart+offset);return;}}",
        "mapping.textContent='';}",
        "document.getElementById('previous').addEventListener('click',()=>goTo(activeIndex-1));",
        "document.getElementById('next').addEventListener('click',()=>goTo(activeIndex+1));",
        "function mapLine(line,isOriginal){const key=isOriginal?'a':'b',other=isOriginal?'b':'a';",
        "const segment=lineMap.find(item=>line>=item[key+'Start']&&line<=Math.max(item[key+'Start'],item[key+'End']));",
        "if(!segment)return 1;const offset=Math.max(0,line-segment[key+'Start']),length=segment[other+'Length'];",
        "return segment[other+'Start']+(length?Math.min(offset,length-1):0);}",
        "function selectObject(){if(!navigationTarget.token||localStorage.getItem(storageKey+':jump')===navigationTarget.token)return;",
        "const beforeSpan=navigationTarget.before,afterSpan=navigationTarget.after;",
        "const aLine=beforeSpan?beforeSpan.start+1:afterSpan?mapLine(afterSpan.start+1,false):1;",
        "const bLine=afterSpan?afterSpan.start+1:beforeSpan?mapLine(beforeSpan.start+1,true):1;",
        "function highlight(editor,span,line){const max=editor.getModel().getLineCount(),start=Math.max(1,Math.min(max,span?span.start+1:line)),",
        "end=Math.max(start,Math.min(max,span?span.end+1:line));editor.revealLineInCenter(start);",
        "const decorations=[];for(let current=start;current<=end;current++)decorations.push({range:new monaco.Range(current,1,current,1),",
        "options:{isWholeLine:true,className:'selected-object-flash'}});const old=editor.deltaDecorations([],decorations);",
        "setTimeout(()=>editor.deltaDecorations(old,[]),2000);}",
        "highlight(left,beforeSpan,aLine);highlight(right,afterSpan,bLine);",
        "if(!beforeSpan||!afterSpan){const editor=beforeSpan?right:left,line=beforeSpan?bLine:aLine,",
        "zone=document.createElement('div');zone.className='missing-policy-marker';zone.textContent='<Policy Missing>';",
        "editor.changeViewZones(accessor=>accessor.addZone({afterLineNumber:Math.max(0,line-1),heightInLines:1,domNode:zone}));}",
        "mapping.textContent='A '+(beforeSpan?beforeSpan.start+1:'Missing')+' -> B '+(afterSpan?afterSpan.start+1:'Missing');",
        "localStorage.setItem(storageKey+':jump',navigationTarget.token);}",
        "function persist(){const state={sync:syncToggle.checked,leftTop:left.getScrollTop(),leftLeft:left.getScrollLeft(),",
        "rightTop:right.getScrollTop(),rightLeft:right.getScrollLeft()};localStorage.setItem(storageKey,JSON.stringify(state));}",
        "function onScroll(source,target,isOriginal){if(restoring)return;persist();",
        "const ranges=source.getVisibleRanges();if(ranges.length)showMapping(ranges[0].startLineNumber,isOriginal);",
        "if(!syncToggle.checked||syncing||!ranges.length)return;syncing=true;",
        "const mapped=mapLine(ranges[0].startLineNumber,isOriginal);target.revealLine(mapped);",
        "requestAnimationFrame(()=>{syncing=false;persist()});}",
        "left.onDidScrollChange(()=>onScroll(left,right,true));right.onDidScrollChange(()=>onScroll(right,left,false));",
        "syncToggle.addEventListener('change',()=>{persist();});",
        "requestAnimationFrame(()=>{if(saved.leftTop!==undefined){left.setScrollTop(saved.leftTop);left.setScrollLeft(saved.leftLeft||0);}",
        "if(saved.rightTop!==undefined){right.setScrollTop(saved.rightTop);right.setScrollLeft(saved.rightLeft||0);}",
        "restoring=false;",
        (
            f"if(navigationToken&&localStorage.getItem(storageKey+':diff-jump')!==navigationToken){{"
            f"goTo({initial_index});localStorage.setItem(storageKey+':diff-jump',navigationToken);}}"
            if focus_row is not None
            else ""
        ),
        "selectObject();",
        "if(!blocks.length)counter.textContent='No differences';else if(activeIndex<0)counter.textContent='Difference 0 of '+blocks.length;});",
        "window.addEventListener('resize',()=>diff.layout());});}",
        "</script></body></html>",
    ]
    return "".join(output)